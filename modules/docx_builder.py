"""
DOCX builder module for creating Word documents from OCR data.

This module handles the construction of Word documents with precise
formatting based on OCR results.
"""

from pathlib import Path
from typing import Dict, Optional, List
from loguru import logger

try:
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    PYTHON_DOCX_AVAILABLE = True
except ImportError:
    PYTHON_DOCX_AVAILABLE = False
    logger.warning("python-docx not installed. Document creation will not be available.")

from config import Config
from models.formatting_models import (
    OCRDocument,
    OCRPage,
    OCRParagraph,
    OCRLine,
    OCRWord,
    CharacterFormatting,
    ParagraphFormatting,
    Alignment,
    Color,
)


class DocxBuilder:
    """
    Builds Word documents from OCR data with formatting.
    """

    def __init__(
        self,
        optimize: bool = Config.DOCX_OPTIMIZATION,
        create_styles: bool = Config.CREATE_STYLES,
    ):
        """
        Initialize DOCX builder.

        Args:
            optimize: Whether to optimize the document
            create_styles: Whether to create reusable styles
        """
        if not PYTHON_DOCX_AVAILABLE:
            raise ImportError(
                "python-docx is required. Install it with: pip install python-docx"
            )

        self.optimize = optimize
        self.create_styles = create_styles

        # Track created styles
        self.custom_styles: Dict[str, str] = {}
        self.style_counter = 0

        logger.info(
            f"Initialized DOCX builder (optimize={optimize}, "
            f"create_styles={create_styles})"
        )

    def create_document(
        self,
        ocr_document: OCRDocument,
        output_path: Path,
    ) -> None:
        """
        Create a Word document from OCR data.

        Args:
            ocr_document: OCRDocument with all pages
            output_path: Path to save the document
        """
        logger.info(f"Creating DOCX document: {output_path}")

        # Create new document
        doc = Document()

        # Set page margins
        self._set_page_margins(doc)

        # Process each page
        for i, page in enumerate(ocr_document.pages):
            logger.debug(f"Adding page {page.page_number + 1} to document")

            # Add page content
            self._add_page_to_document(doc, page)

            # Add page break (except for last page)
            if Config.PRESERVE_PAGE_BREAKS and i < len(ocr_document.pages) - 1:
                doc.add_page_break()

        # Optimize if enabled
        if self.optimize:
            self._optimize_document(doc)

        # Save document
        doc.save(str(output_path))

        logger.info(f"Document saved: {output_path}")

    def _set_page_margins(self, doc: Document) -> None:
        """
        Set page margins.

        Args:
            doc: Document to modify
        """
        sections = doc.sections
        for section in sections:
            section.top_margin = Inches(Config.MARGIN_TOP)
            section.bottom_margin = Inches(Config.MARGIN_BOTTOM)
            section.left_margin = Inches(Config.MARGIN_LEFT)
            section.right_margin = Inches(Config.MARGIN_RIGHT)

    def _add_page_to_document(self, doc: Document, page: OCRPage) -> None:
        """
        Add a page to the document.

        Args:
            doc: Document to add to
            page: OCRPage to add
        """
        for paragraph in page.paragraphs:
            self._add_paragraph_to_document(doc, paragraph)

    def _add_paragraph_to_document(
        self,
        doc: Document,
        ocr_paragraph: OCRParagraph,
    ) -> None:
        """
        Add a paragraph to the document.

        Args:
            doc: Document to add to
            ocr_paragraph: OCRParagraph to add
        """
        if not ocr_paragraph.lines:
            return

        # Create paragraph
        para = doc.add_paragraph()

        # Apply paragraph formatting
        self._apply_paragraph_formatting(para, ocr_paragraph.para_formatting)

        # Add all words from all lines
        for line in ocr_paragraph.lines:
            for word in line.words:
                # Add word with formatting
                run = para.add_run(word.text + ' ')
                self._apply_character_formatting(run, word.char_formatting)

    def _apply_paragraph_formatting(
        self,
        para,
        formatting: ParagraphFormatting,
    ) -> None:
        """
        Apply paragraph-level formatting.

        Args:
            para: Word paragraph object
            formatting: ParagraphFormatting data
        """
        # Alignment
        alignment_map = {
            Alignment.LEFT: WD_ALIGN_PARAGRAPH.LEFT,
            Alignment.CENTER: WD_ALIGN_PARAGRAPH.CENTER,
            Alignment.RIGHT: WD_ALIGN_PARAGRAPH.RIGHT,
            Alignment.JUSTIFIED: WD_ALIGN_PARAGRAPH.JUSTIFY,
        }
        para.alignment = alignment_map.get(formatting.alignment, WD_ALIGN_PARAGRAPH.LEFT)

        # Paragraph format
        para_format = para.paragraph_format

        # Indentation
        if formatting.indent_first_line:
            para_format.first_line_indent = Pt(formatting.indent_first_line)

        if formatting.indent_left:
            para_format.left_indent = Pt(formatting.indent_left)

        if formatting.indent_right:
            para_format.right_indent = Pt(formatting.indent_right)

        # Spacing
        if formatting.spacing_before:
            para_format.space_before = Pt(formatting.spacing_before)

        if formatting.spacing_after:
            para_format.space_after = Pt(formatting.spacing_after)

        # Line spacing
        if formatting.line_spacing and formatting.line_spacing != 1.0:
            para_format.line_spacing = formatting.line_spacing

        # Keep with next
        if formatting.keep_with_next:
            para_format.keep_with_next = True

        # Keep together
        if formatting.keep_together:
            para_format.keep_together = True

        # Page break before
        if formatting.page_break_before:
            para_format.page_break_before = True

    def _apply_character_formatting(
        self,
        run,
        formatting: CharacterFormatting,
    ) -> None:
        """
        Apply character-level formatting.

        Args:
            run: Word run object
            formatting: CharacterFormatting data
        """
        # Font
        font = run.font

        # Font family
        if formatting.font_family:
            font.name = formatting.font_family

        # Font size
        if formatting.font_size:
            font.size = Pt(formatting.font_size)

        # Bold
        if formatting.is_bold:
            font.bold = True

        # Italic
        if formatting.is_italic:
            font.italic = True

        # Underline
        if formatting.is_underlined:
            font.underline = True

        # Strikethrough
        if formatting.is_strikethrough:
            font.strike = True

        # Small caps
        if formatting.is_small_caps:
            font.small_caps = True

        # Superscript
        if formatting.is_superscript:
            font.superscript = True

        # Subscript
        if formatting.is_subscript:
            font.subscript = True

        # Text color
        if formatting.text_color and not formatting.text_color.is_black():
            color = formatting.text_color
            font.color.rgb = RGBColor(color.r, color.g, color.b)

        # Background color (highlight)
        if formatting.background_color and not formatting.background_color.is_white():
            # Note: python-docx has limited highlight support
            # We'll skip background colors for now
            pass

    def _optimize_document(self, doc: Document) -> None:
        """
        Optimize the document.

        Args:
            doc: Document to optimize
        """
        logger.debug("Optimizing document")

        # Merge similar paragraphs, remove empty paragraphs, etc.
        # This is a placeholder for optimization logic

        # Remove empty paragraphs
        paragraphs_to_remove = []
        for para in doc.paragraphs:
            if not para.text.strip():
                paragraphs_to_remove.append(para)

        for para in paragraphs_to_remove:
            # Removing paragraphs in python-docx is tricky
            # We'll skip this for now
            pass

    def create_style(
        self,
        doc: Document,
        style_name: str,
        char_formatting: CharacterFormatting,
        para_formatting: ParagraphFormatting,
    ) -> str:
        """
        Create a reusable style.

        Args:
            doc: Document to add style to
            style_name: Name for the style
            char_formatting: Character formatting
            para_formatting: Paragraph formatting

        Returns:
            Style name
        """
        # This is a simplified implementation
        # Full implementation would create actual Word styles
        self.style_counter += 1
        style_id = f"CustomStyle{self.style_counter}"
        self.custom_styles[style_name] = style_id

        return style_id


class DocumentValidator:
    """
    Validates Word documents.
    """

    @staticmethod
    def validate_document(doc_path: Path) -> bool:
        """
        Validate a Word document.

        Args:
            doc_path: Path to document

        Returns:
            True if valid
        """
        try:
            doc = Document(str(doc_path))

            # Basic validation
            if not doc.paragraphs:
                logger.warning("Document has no paragraphs")
                return False

            logger.info(f"Document is valid with {len(doc.paragraphs)} paragraphs")
            return True

        except Exception as e:
            logger.error(f"Document validation failed: {e}")
            return False


class DocumentMetadata:
    """
    Manages document metadata.
    """

    @staticmethod
    def set_metadata(
        doc: Document,
        title: Optional[str] = None,
        author: Optional[str] = None,
        subject: Optional[str] = None,
        keywords: Optional[str] = None,
    ) -> None:
        """
        Set document metadata.

        Args:
            doc: Document to modify
            title: Document title
            author: Document author
            subject: Document subject
            keywords: Document keywords
        """
        core_properties = doc.core_properties

        if title:
            core_properties.title = title

        if author:
            core_properties.author = author

        if subject:
            core_properties.subject = subject

        if keywords:
            core_properties.keywords = keywords

        logger.debug("Set document metadata")


def create_docx_builder(
    optimize: bool = Config.DOCX_OPTIMIZATION,
) -> DocxBuilder:
    """
    Factory function to create a DocxBuilder.

    Args:
        optimize: Whether to optimize documents

    Returns:
        DocxBuilder instance
    """
    return DocxBuilder(optimize=optimize)


def create_simple_document(text: str, output_path: Path) -> None:
    """
    Create a simple document from plain text.

    Args:
        text: Text content
        output_path: Path to save document
    """
    doc = Document()

    for para_text in text.split('\n\n'):
        if para_text.strip():
            doc.add_paragraph(para_text.strip())

    doc.save(str(output_path))
    logger.info(f"Created simple document: {output_path}")
